"""Build the phoning prospect workbook from parsed.json."""
import json, re, sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as CL

SRC, OUT = sys.argv[1], sys.argv[2]
data = json.load(open(SRC))
# A company name ending with ',' continues on the next (regular-weight) line in the PDF.
for _p in data:
    for _c in _p['contacts']:
        if _c['societe'].endswith(',') and _c['adresse']:
            _c['societe'] = _c['societe'] + ' ' + _c['adresse'].pop(0)

# ------------------------------------------------------------------ translations
STADE = {'Baubewilligung erteilt': 'Permis délivré', 'Baugesuch eingereicht': 'Demande de permis déposée'}
EVENT = {'Baubewilligung erteilt': 'Permis délivré', 'Baugesuch publiziert': 'Demande publiée'}
ROLE = {'Bauherr': "Maître d'ouvrage", 'Architekt/Planer': 'Architecte / planificateur',
        'Generalunternehmer': 'Entreprise générale', 'Bauingenieur': 'Ingénieur civil',
        'Bauherr / Architekt/Planer': "Maître d'ouvrage & architecte"}
ROLE_ORDER = {"Maître d'ouvrage": 0, "Maître d'ouvrage & architecte": 0, 'Architecte / planificateur': 1,
              'Entreprise générale': 2, 'Ingénieur civil': 3}
KAT = {'Handel und Verwaltung': 'Commerce et administration',
       'Unterricht, Bildung und Forschung': 'Enseignement, formation et recherche',
       'Industrie und Gewerbe': 'Industrie et artisanat', 'Land- und Forstwirtschaft': 'Agriculture et sylviculture',
       'Kultur und Geselligkeit': 'Culture et loisirs', 'Fürsorge und Gesundheit': 'Social et santé',
       'Militär- und Schutzanlagen': 'Installations militaires et de protection'}
ZWECK = {'Nicht relevant': 'Non pertinent', 'Eigenbedarf': 'Usage propre'}
WORK = {'Umbau innen': 'transformation intérieure', 'Umbau aussen': 'transformation extérieure',
        'Neubau': 'construction neuve', 'Anbau': 'agrandissement', 'Abbruch': 'démolition'}
TYPE = {
    'Restaurationsbetriebe': 'Restaurants', 'Wohnungen': 'Logements',
    'Aussenanlagen, Kinderspielplätze und Parkanlagen': 'Aménagements extérieurs, places de jeux et parcs',
    'Fitnesscenter/-raum': 'Centre / salle de fitness', 'Parkhäuser und Einstellhallen': 'Parkings et garages souterrains',
    'Garagen / Fertiggaragen': 'Garages / garages préfabriqués', 'Feuerwehrgebäude': 'Casernes de pompiers',
    'Konzertbauten und Theaterbauten': 'Salles de concert et théâtres', 'Strassen': 'Routes',
    'Primar- und Sekundarschulen': 'Écoles primaires et secondaires',
    'Berufs- und höhere Fachschulen': 'Écoles professionnelles et hautes écoles spécialisées',
    'Mittelschulen und Gymnasien': 'Écoles moyennes et gymnases', 'Hochschulen und Universitäten': 'Hautes écoles et universités',
    'Forschungsinstitute': 'Instituts de recherche', 'Lagerhallen': 'Halles de stockage', 'Industriehallen': 'Halles industrielles',
    'Industrielle Produktionsbauten': 'Bâtiments de production industrielle',
    'Betriebs- und Gewerbebauten': "Bâtiments d'exploitation et artisanaux", 'Atelier und Studio': 'Ateliers et studios',
    'Schuppen und Hütten': 'Hangars et cabanes', 'Futterlagerräume, Treibhäuser und Silobauten': 'Fourragères, serres et silos',
    'Stallungen und landwirtschaftliche Produktionsanlagen': 'Étables et installations agricoles',
    'Tierspitäler': 'Cliniques vétérinaires', 'Jauchegrube': 'Fosses à purin', 'Ladenbauten': 'Commerces / magasins',
    'Bürobauten mit einfachen Anforderungen': 'Bureaux (standard simple)',
    'Bürobauten mit erhöhten Anforderungen': 'Bureaux (standard élevé)',
    'Verwaltungsgebäude und Rechenzentren': 'Bâtiments administratifs et centres de calcul',
    'Banken, Postgebäude und Fernmeldegebäude': 'Banques, postes et télécommunications',
    'Gemeindehäuser, Rathäuser und Regierungsgebäude': 'Maisons communales et bâtiments gouvernementaux',
    'Wiedereingliederungsstätten': 'Centres de réinsertion', 'Arztpraxen und Ärztehäuser': 'Cabinets médicaux',
}


def tr_sub(s):
    m = re.match(r'^(\d+)\s+(.+?)\s+\((.+)\)$', s)
    if not m:
        return s
    code, typ, work = m.groups()
    return f"{TYPE.get(typ, typ)} – {WORK.get(work, work)}"


def civ(name):
    return re.sub(r'^Herr\b', 'M.', re.sub(r'^Frau\b', 'Mme', name))


def to_date(s):
    return datetime.strptime(s, '%d.%m.%Y') if s else None


def montant(s):
    m = re.match(r'([\d.]+)\s*Mio', s or '')
    return round(float(m.group(1)) * 1_000_000) if m else None


def tel_link(t):
    digits = re.sub(r'\D', '', t)
    return 'tel:+41' + digits[1:] if digits.startswith('0') else None


def nat_key(name):
    return [int(p) if p.isdigit() else p.lower() for p in re.split(r'(\d+)', name)]


# ------------------------------------------------------------------ duplicates (same Infopro object in 2 files)
by_obj = {}
for p in data:
    by_obj.setdefault(p['kv']['Objektnummer'], []).append(f"{p['fichier']} (p.{p['page']})")


def dup_of(p):
    others = [x for x in by_obj[p['kv']['Objektnummer']] if x != f"{p['fichier']} (p.{p['page']})"]
    return 'Même projet que ' + ', '.join(others) if others else ''


# ------------------------------------------------------------------ styles
F = 'Arial'
font = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
hfont = Font(name=F, size=10, bold=True, color='FFFFFF')
title_font = Font(name=F, size=16, bold=True, color='0E3A40')
h2 = Font(name=F, size=12, bold=True, color='0E3A40')
muted = Font(name=F, size=9, italic=True, color='5C6B6D')
link = Font(name=F, size=10, color='1C5F68', underline='single')
missing = Font(name=F, size=10, italic=True, color='B4541A')
FILL_H = PatternFill('solid', fgColor='0E3A40')      # data header
FILL_H_IN = PatternFill('solid', fgColor='B4541A')   # header of columns to fill in
FILL_H_DET = PatternFill('solid', fgColor='5C6B6D')  # header of detail columns
FILL_IN = PatternFill('solid', fgColor='FFF2CC')     # cells to fill in (yellow)
FILL_SOFT = PatternFill('solid', fgColor='EAF2F2')
thin = Side(style='thin', color='DEE5E3')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')
CENTER = Alignment(horizontal='center', vertical='top')
DATE_FMT = 'dd/mm/yyyy'
CHF_FMT = "#,##0 \"CHF\""

wb = Workbook()

# ------------------------------------------------------------------ Listes (dropdowns + parameters)
ls = wb.active
ls.title = 'Listes'
LISTS = {
    'Résultat appel': ['Joint – bon interlocuteur', 'Joint – barrage / standard', 'Messagerie (message laissé)',
                       'Messagerie (sans message)', 'Pas de réponse', 'Occupé', 'Numéro erroné / inexistant'],
    'Qualification': ['À appeler', 'À rappeler', 'Intéressé', 'RDV fixé', 'Offre demandée', 'Pas intéressé',
                      'Déjà équipé / fournisseur en place', 'Hors cible', 'Projet terminé / annulé'],
    'Prochaine action': ['Rappeler', 'Envoyer documentation', 'Envoyer offre', 'Rendez-vous / visite',
                         'Relance e-mail', 'Aucune (dossier clos)'],
}
for j, (name, vals) in enumerate(LISTS.items()):
    c = 1 + j * 2
    ls.cell(1, c, name).font = bold
    for i, v in enumerate(vals, 2):
        ls.cell(i, c, v).font = font
    ls.column_dimensions[CL(c)].width = 34
    ls.column_dimensions[CL(c + 1)].width = 3
# priority thresholds
ls['H1'] = 'Paramètres de priorité'; ls['H1'].font = bold
ls['H2'] = 'Seuil priorité A (CHF ≥)'; ls['I2'] = 300000
ls['H3'] = 'Seuil priorité B (CHF ≥)'; ls['I3'] = 100000
ls['H4'] = 'En dessous du seuil B → priorité C'
for r in (2, 3, 4):
    ls[f'H{r}'].font = font
for r in (2, 3):
    ls[f'I{r}'].font = Font(name=F, size=10, color='0000FF')
    ls[f'I{r}'].fill = FILL_IN
    ls[f'I{r}'].number_format = CHF_FMT
ls['H6'] = ('Seuils par défaut proposés (à ajuster) : montant des travaux annoncé dans la fiche Infopro. '
            'Modifier les cellules jaunes recalcule la colonne « Priorité » de l’onglet Appels.')
ls['H6'].font = muted
ls['H6'].alignment = Alignment(wrap_text=True, vertical='top')
ls.merge_cells('H6:I9')
ls.column_dimensions['H'].width = 34
ls.column_dimensions['I'].width = 16


def list_ref(name):
    j = list(LISTS).index(name)
    col = CL(1 + j * 2)
    return f"Listes!${col}$2:${col}${len(LISTS[name]) + 1}"


# ------------------------------------------------------------------ Appels (one row per contact × project)
ws = wb.create_sheet('Appels', 0)
COLS = [
    # (header, width, group)  group: core | input | detail
    ('Fichier PDF', 21, 'core'), ('Page', 6, 'core'), ('Priorité', 8, 'core'), ('Rôle', 17, 'core'),
    ('Société', 34, 'core'), ('Interlocuteur', 20, 'core'), ('Téléphone', 15, 'core'), ('E-mail', 26, 'core'),
    ('Projet', 34, 'core'), ('Lieu du chantier', 18, 'core'), ('Stade', 16, 'core'),
    ('Montant travaux (CHF)', 14, 'core'), ('Projets liés (même contact)', 10, 'core'),
    ('Résultat dernier appel', 22, 'input'), ('Qualification', 20, 'input'), ('Nb tentatives', 9, 'input'),
    ('Date dernier appel', 12, 'input'), ('Prochaine action', 19, 'input'), ('Date de relance', 12, 'input'),
    ('Commentaires / interlocuteur joint', 40, 'input'),
    ('Réf. projet Infopro', 11, 'detail'), ('Date du stade', 11, 'detail'), ('Adresse du contact', 30, 'detail'),
    ('Site web', 22, 'detail'), ('Adresse du chantier', 28, 'detail'), ('Catégorie', 22, 'detail'),
    ('Description du projet', 50, 'detail'), ('Source du contact', 14, 'detail'), ('Doublon', 24, 'detail'),
]
H = {name: i + 1 for i, (name, _, _) in enumerate(COLS)}
for (name, width, grp), i in zip(COLS, range(1, len(COLS) + 1)):
    c = ws.cell(1, i, name)
    c.font = hfont
    c.fill = {'core': FILL_H, 'input': FILL_H_IN, 'detail': FILL_H_DET}[grp]
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    c.border = BORDER
    ws.column_dimensions[CL(i)].width = width
ws.row_dimensions[1].height = 42

rows = []
for p in data:
    last = p['termine'][-1] if p['termine'] else {'date': '', 'evenement': ''}
    for k, c in enumerate(p['contacts']):
        rows.append((p, c, k))
rows.sort(key=lambda t: (nat_key(t[0]['fichier']), t[0]['page'], ROLE_ORDER.get(ROLE.get(t[1]['role'], ''), 9), t[2]))

r = 1
for p, c, _ in rows:
    r += 1
    last = p['termine'][-1] if p['termine'] else {'date': ''}
    role = ROLE.get(c['role'], c['role'])
    tel = c['tel']
    vals = {
        'Fichier PDF': p['fichier'], 'Page': p['page'], 'Rôle': role, 'Société': c['societe'],
        'Interlocuteur': ', '.join(civ(x) for x in c['interlocuteurs']),
        'Téléphone': tel or 'À rechercher', 'E-mail': c['email'],
        'Projet': p['projet'], 'Lieu du chantier': p['lieu_titre'], 'Stade': STADE.get(p['stade'], p['stade']),
        'Montant travaux (CHF)': montant(p['kv'].get('Bausumme')),
        'Qualification': 'À appeler', 'Nb tentatives': 0,
        'Réf. projet Infopro': p['kv'].get('Objektnummer'), 'Date du stade': to_date(last['date']),
        'Adresse du contact': ', '.join(c['adresse']), 'Site web': c['web'],
        'Adresse du chantier': p['adresse_chantier'], 'Catégorie': KAT.get(p['kv'].get('Kategorie'), p['kv'].get('Kategorie')),
        'Description du projet': p['description'], 'Source du contact': c['source'], 'Doublon': dup_of(p),
    }
    for name, v in vals.items():
        ws.cell(r, H[name], v)
    # formulas
    cm, cp, ct, cs = CL(H['Montant travaux (CHF)']), CL(H['Priorité']), CL(H['Téléphone']), CL(H['Société'])
    ws.cell(r, H['Priorité'], f'=IF({cm}{r}="","C",IF({cm}{r}>=Listes!$I$2,"A",IF({cm}{r}>=Listes!$I$3,"B","C")))')
    ws.cell(r, H['Projets liés (même contact)'],
            f'=IF({ct}{r}="À rechercher",COUNTIF(${cs}$2:${cs}$9999,{cs}{r}),COUNTIF(${ct}$2:${ct}$9999,{ct}{r}))')
    # hyperlinks
    if tel and tel_link(tel):
        ws.cell(r, H['Téléphone']).hyperlink = tel_link(tel)
    if c['email']:
        ws.cell(r, H['E-mail']).hyperlink = 'mailto:' + c['email']
    if c['web']:
        ws.cell(r, H['Site web']).hyperlink = c['web']
LAST = r

# cell styling
for row in ws.iter_rows(min_row=2, max_row=LAST):
    for cell in row:
        name = COLS[cell.column - 1][0]
        grp = COLS[cell.column - 1][2]
        cell.font = font
        cell.border = BORDER
        cell.alignment = WRAP if name in ('Projet', 'Société', 'Commentaires / interlocuteur joint', 'Rôle') else TOP
        if grp == 'input':
            cell.fill = FILL_IN
        if name in ('Page', 'Priorité', 'Nb tentatives', 'Projets liés (même contact)'):
            cell.alignment = CENTER
        if name in ('Date dernier appel', 'Date de relance', 'Date du stade'):
            cell.number_format = DATE_FMT
        if name == 'Montant travaux (CHF)':
            cell.number_format = CHF_FMT
        if name == 'Priorité':
            cell.font = bold
        if name in ('Téléphone', 'E-mail', 'Site web') and cell.hyperlink:
            cell.font = link
        if name == 'Téléphone' and cell.value == 'À rechercher':
            cell.font = missing
        if name == 'Société':
            cell.font = bold

# table, freeze, grouping
ref = f"A1:{CL(len(COLS))}{LAST}"
tbl = Table(displayName='tblAppels', ref=ref)
tbl.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
ws.add_table(tbl)
ws.freeze_panes = 'F2'
first_det, last_det = H['Réf. projet Infopro'], len(COLS)
for ci in range(first_det, last_det + 1):  # per-column outline keeps each column's own width
    ws.column_dimensions[CL(ci)].outlineLevel = 1
for rr in range(2, LAST + 1):
    ws.row_dimensions[rr].height = 30  # fixed 2-line rows: the call sheet stays scannable

# data validation
def add_dv(col_name, formula):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True,
                        errorTitle='Valeur non prévue', error='Choisissez une valeur dans la liste (onglet Listes).')
    ws.add_data_validation(dv)
    col = CL(H[col_name])
    dv.add(f"{col}2:{col}{LAST + 500}")

add_dv('Résultat dernier appel', list_ref('Résultat appel'))
add_dv('Qualification', list_ref('Qualification'))
add_dv('Prochaine action', list_ref('Prochaine action'))
for nm in ('Date dernier appel', 'Date de relance'):
    dv = DataValidation(type='date', operator='greaterThan', formula1='DATE(2020,1,1)', allow_blank=True,
                        showErrorMessage=True, errorTitle='Date invalide', error='Saisissez une date (jj/mm/aaaa).')
    ws.add_data_validation(dv)
    dv.add(f"{CL(H[nm])}2:{CL(H[nm])}{LAST + 500}")
dv = DataValidation(type='whole', operator='between', formula1='0', formula2='50', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"{CL(H['Nb tentatives'])}2:{CL(H['Nb tentatives'])}{LAST + 500}")

# conditional formatting
q, rl, pr = CL(H['Qualification']), CL(H['Date de relance']), CL(H['Priorité'])
rng_all = f"A2:{CL(len(COLS))}{LAST + 500}"
# whole row tint by qualification (soft), keyed on the qualification column
ws.conditional_formatting.add(rng_all, FormulaRule(formula=[f'OR(${q}2="RDV fixé",${q}2="Offre demandée")'],
                              fill=PatternFill('solid', fgColor='C6EFCE')))
ws.conditional_formatting.add(rng_all, FormulaRule(formula=[f'${q}2="Intéressé"'],
                              fill=PatternFill('solid', fgColor='E2F0D9')))
ws.conditional_formatting.add(rng_all, FormulaRule(
    formula=[f'OR(${q}2="Pas intéressé",${q}2="Hors cible",${q}2="Projet terminé / annulé",${q}2="Déjà équipé / fournisseur en place")'],
    font=Font(color='8C8C8C'), fill=PatternFill('solid', fgColor='F2F2F2')))
# overdue follow-up (date passed and file still open)
ws.conditional_formatting.add(f"{rl}2:{rl}{LAST + 500}", FormulaRule(
    formula=[f'AND({rl}2<>"",{rl}2<=TODAY(),OR(${q}2="À rappeler",${q}2="Intéressé",${q}2="Offre demandée",${q}2="À appeler"))'],
    font=Font(name=F, bold=True, color='9C0006'), fill=PatternFill('solid', fgColor='FFC7CE')))
ws.conditional_formatting.add(f"{pr}2:{pr}{LAST + 500}", CellIsRule(operator='equal', formula=['"A"'],
                              font=Font(name=F, bold=True, color='FFFFFF'), fill=PatternFill('solid', fgColor='B4541A')))
ws.conditional_formatting.add(f"{pr}2:{pr}{LAST + 500}", CellIsRule(operator='equal', formula=['"B"'],
                              fill=PatternFill('solid', fgColor='FCE4D6')))

# header comments (guidance)
notes = {
    'Priorité': 'Calculée selon le montant des travaux (seuils réglables dans l’onglet Listes).',
    'Projets liés (même contact)': 'Nombre de lignes avec le même numéro (ou la même société si pas de numéro). '
                                   'Si > 1 : un seul appel peut couvrir plusieurs projets.',
    'Résultat dernier appel': 'Ce qui s’est passé techniquement lors de l’appel (joint, messagerie…).',
    'Qualification': 'Où en est le prospect commercialement. « À appeler » par défaut.',
    'Date de relance': 'Toujours renseigner une date de relance : passe en rouge quand elle est échue.',
    'Téléphone': 'Cliquable : lance l’appel via votre softphone / Teams / Skype (lien tel:+41…).',
    'Source du contact': '« fiche » = bloc de contact détaillé ; « liste complémentaire » = tableau résumé en bas de fiche.',
}
for k, v in notes.items():
    ws.cell(1, H[k]).comment = Comment(v, 'Prospection')

# ------------------------------------------------------------------ Projets (one row per project)
wp = wb.create_sheet('Projets', 1)
PCOLS = [('Fichier PDF', 21), ('Page', 6), ('Réf. projet Infopro', 11), ('Projet', 38), ('Lieu du chantier', 18),
         ('Adresse du chantier', 28), ('Parcelle', 12), ('Stade', 18), ('Demande publiée le', 12),
         ('Permis délivré le', 12), ('Montant travaux (CHF)', 14), ('Catégorie', 24), ('Affectation', 13),
         ('Sous-catégories', 42), ('Description', 55), ('Date de recherche Infopro', 12),
         ("Maître d'ouvrage", 28), ('Architecte / planificateur', 28), ('Nb contacts', 9),
         ('Contacts joints', 9), ('Intéressés / RDV / offres', 10), ('Doublon', 26)]
PH = {n: i + 1 for i, (n, _) in enumerate(PCOLS)}
for i, (n, w) in enumerate(PCOLS, 1):
    c = wp.cell(1, i, n)
    c.font, c.fill, c.border = hfont, FILL_H, BORDER
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    wp.column_dimensions[CL(i)].width = w
wp.row_dimensions[1].height = 42
A = {n: CL(i) for n, i in H.items()}  # column letters in Appels
projects = sorted(data, key=lambda p: (nat_key(p['fichier']), p['page']))
for r, p in enumerate(projects, 2):
    ev = {EVENT.get(t['evenement'], t['evenement']): to_date(t['date']) for t in p['termine']}
    mo = [c['societe'] for c in p['contacts'] if c['role'].startswith('Bauherr') and c['source'] == 'fiche']
    ar = [c['societe'] for c in p['contacts'] if 'Architekt' in c['role'] and c['source'] == 'fiche']
    vals = {'Fichier PDF': p['fichier'], 'Page': p['page'], 'Réf. projet Infopro': p['kv'].get('Objektnummer'),
            'Projet': p['projet'], 'Lieu du chantier': p['lieu_titre'], 'Adresse du chantier': p['adresse_chantier'],
            'Parcelle': p['reference'], 'Stade': STADE.get(p['stade'], p['stade']),
            'Demande publiée le': ev.get('Demande publiée'), 'Permis délivré le': ev.get('Permis délivré'),
            'Montant travaux (CHF)': montant(p['kv'].get('Bausumme')),
            'Catégorie': KAT.get(p['kv'].get('Kategorie'), p['kv'].get('Kategorie')),
            'Affectation': ZWECK.get(p['kv'].get('Verwendungszweck'), p['kv'].get('Verwendungszweck') or ''),
            'Sous-catégories': '\n'.join(tr_sub(s) for s in p['sous_categories']),
            'Description': p['description'], 'Date de recherche Infopro': to_date(p['kv'].get('Recherchedatum')),
            "Maître d'ouvrage": ' / '.join(mo), 'Architecte / planificateur': ' / '.join(ar), 'Doublon': dup_of(p)}
    for n, v in vals.items():
        wp.cell(r, PH[n], v)
    key = f'$A{r}'
    f_pdf, f_pg = f"Appels!${A['Fichier PDF']}$2:${A['Fichier PDF']}$9999", f"Appels!${A['Page']}$2:${A['Page']}$9999"
    f_res, f_q = (f"Appels!${A['Résultat dernier appel']}$2:${A['Résultat dernier appel']}$9999",
                  f"Appels!${A['Qualification']}$2:${A['Qualification']}$9999")
    wp.cell(r, PH['Nb contacts'], f'=COUNTIFS({f_pdf},{key},{f_pg},$B{r})')
    wp.cell(r, PH['Contacts joints'],
            f'=COUNTIFS({f_pdf},{key},{f_pg},$B{r},{f_res},"Joint*")')
    wp.cell(r, PH['Intéressés / RDV / offres'],
            f'=COUNTIFS({f_pdf},{key},{f_pg},$B{r},{f_q},"Intéressé")'
            f'+COUNTIFS({f_pdf},{key},{f_pg},$B{r},{f_q},"RDV fixé")'
            f'+COUNTIFS({f_pdf},{key},{f_pg},$B{r},{f_q},"Offre demandée")')
PLAST = len(projects) + 1
for row in wp.iter_rows(min_row=2, max_row=PLAST):
    for cell in row:
        n = PCOLS[cell.column - 1][0]
        cell.font, cell.border = font, BORDER
        cell.alignment = WRAP if n in ('Projet', 'Description', 'Sous-catégories', 'Adresse du chantier', "Maître d'ouvrage",
                                       'Architecte / planificateur', 'Doublon', 'Catégorie') else TOP
        if n in ('Demande publiée le', 'Permis délivré le', 'Date de recherche Infopro'):
            cell.number_format = DATE_FMT
        if n == 'Montant travaux (CHF)':
            cell.number_format = CHF_FMT
        if n in ('Page', 'Nb contacts', 'Contacts joints', 'Intéressés / RDV / offres'):
            cell.alignment = CENTER
        if n == 'Projet':
            cell.font = bold
t2 = Table(displayName='tblProjets', ref=f"A1:{CL(len(PCOLS))}{PLAST}")
t2.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
wp.add_table(t2)
wp.freeze_panes = 'E2'
wp.conditional_formatting.add(f"V2:V{PLAST}", FormulaRule(formula=['V2<>""'], fill=PatternFill('solid', fgColor='FCE4D6')))
wp.conditional_formatting.add(f"{CL(PH['Intéressés / RDV / offres'])}2:{CL(PH['Intéressés / RDV / offres'])}{PLAST}",
                              CellIsRule(operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='C6EFCE'),
                                         font=Font(name=F, bold=True)))

# ------------------------------------------------------------------ Tableau de bord
wd = wb.create_sheet('Tableau de bord', 0)
wd.sheet_view.showGridLines = False
wd.column_dimensions['A'].width = 3
wd.column_dimensions['B'].width = 38
wd.column_dimensions['C'].width = 14
wd.column_dimensions['D'].width = 12
wd.column_dimensions['E'].width = 4
wd.column_dimensions['F'].width = 36
wd.column_dimensions['G'].width = 12
wd.column_dimensions['H'].width = 12
wd['B2'] = 'Prospection téléphonique — projets de construction (fiches Infopro « DRINGEND »)'
wd['B2'].font = title_font
wd['B3'] = (f"{len(set(p['fichier'] for p in data))} fichiers PDF · {len(data)} projets · {LAST - 1} contacts. "
            "Les chiffres ci-dessous se mettent à jour automatiquement à partir de l’onglet « Appels ».")
wd['B3'].font = muted

ap = lambda n: f"Appels!${A[n]}$2:${A[n]}$9999"
def kpi(row, label, formula, fmt=None, note=None):
    wd.cell(row, 2, label).font = font
    c = wd.cell(row, 3, formula)
    c.font = Font(name=F, size=11, bold=True, color='0E3A40')
    c.alignment = Alignment(horizontal='right')
    if fmt:
        c.number_format = fmt
    if note:
        wd.cell(row, 4, note).font = muted
    for col in (2, 3):
        wd.cell(row, col).border = Border(bottom=thin)

wd['B5'] = 'Vue d’ensemble'; wd['B5'].font = h2
kpi(6, 'Projets', f"=COUNTA(Projets!$A$2:$A${PLAST})")
kpi(7, 'Contacts à appeler', f"=COUNTA({ap('Société')})")
kpi(8, 'Contacts sans téléphone (à rechercher)', f'=COUNTIF({ap("Téléphone")},"À rechercher")')
kpi(9, 'Montant cumulé des travaux', f"=SUM(Projets!$K$2:$K${PLAST})", CHF_FMT, 'doublons inclus')
wd['B11'] = 'Avancement des appels'; wd['B11'].font = h2
kpi(12, 'Contacts déjà appelés (≥ 1 tentative)', f'=COUNTIF({ap("Nb tentatives")},">0")')
kpi(13, 'Taux de couverture', '=IFERROR(C12/C7,0)', '0%')
kpi(14, 'Total des tentatives', f'=SUM({ap("Nb tentatives")})')
kpi(15, 'Contacts joints', f'=COUNTIF({ap("Résultat dernier appel")},"Joint*")')
kpi(16, 'Taux de joignabilité', '=IFERROR(C15/C12,0)', '0%', 'joints / appelés')
kpi(17, 'Intéressés + RDV + offres', f'=COUNTIF({ap("Qualification")},"Intéressé")+COUNTIF({ap("Qualification")},"RDV fixé")+COUNTIF({ap("Qualification")},"Offre demandée")')
kpi(18, 'Taux de conversion', '=IFERROR(C17/C15,0)', '0%', 'qualifiés / joints')
kpi(19, 'Relances échues (à traiter)',
    f'=SUMPRODUCT(({ap("Date de relance")}<>"")*({ap("Date de relance")}<=TODAY())*'
    f'(({ap("Qualification")}="À rappeler")+({ap("Qualification")}="Intéressé")+({ap("Qualification")}="Offre demandée")+({ap("Qualification")}="À appeler")))')
wd['C19'].font = Font(name=F, size=11, bold=True, color='9C0006')

wd['F5'] = 'Répartition par qualification'; wd['F5'].font = h2
wd['F6'], wd['G6'] = 'Qualification', 'Contacts'
for c in ('F6', 'G6'):
    wd[c].font, wd[c].fill = hfont, FILL_H
for i, v in enumerate(LISTS['Qualification'], 7):
    wd.cell(i, 6, v).font = font
    wd.cell(i, 7, f'=COUNTIF({ap("Qualification")},F{i})').font = font
    for col in (6, 7):
        wd.cell(i, col).border = Border(bottom=thin)
end_q = 6 + len(LISTS['Qualification'])
wd.cell(end_q + 1, 6, 'Total').font = bold
wd.cell(end_q + 1, 7, f'=SUM(G7:G{end_q})').font = bold

row0 = end_q + 3
wd.cell(row0, 6, 'Répartition par priorité').font = h2
wd.cell(row0 + 1, 6, 'Priorité').font = hfont; wd.cell(row0 + 1, 6).fill = FILL_H
wd.cell(row0 + 1, 7, 'Contacts').font = hfont; wd.cell(row0 + 1, 7).fill = FILL_H
wd.cell(row0 + 1, 8, 'Restant à appeler').font = hfont; wd.cell(row0 + 1, 8).fill = FILL_H
for i, v in enumerate(['A', 'B', 'C'], row0 + 2):
    wd.cell(i, 6, v).font = bold
    wd.cell(i, 7, f'=COUNTIF({ap("Priorité")},F{i})').font = font
    wd.cell(i, 8, f'=COUNTIFS({ap("Priorité")},F{i},{ap("Qualification")},"À appeler")').font = font

row1 = row0 + 6
wd.cell(row1, 6, 'Répartition par stade du projet').font = h2
for col, t in ((6, 'Stade'), (7, 'Projets')):
    wd.cell(row1 + 1, col, t).font = hfont
    wd.cell(row1 + 1, col).fill = FILL_H
for i, v in enumerate(sorted(set(STADE.values())), row1 + 2):
    wd.cell(i, 6, v).font = font
    wd.cell(i, 7, f'=COUNTIF(Projets!$H$2:$H${PLAST},F{i})').font = font

tips_row = row1 + 6
wd.cell(tips_row, 2, 'Rappels').font = h2
tips = [
    '1. Filtrez l’onglet « Appels » sur Qualification = « À appeler » et triez par Priorité.',
    '2. Après chaque appel, remplissez tout de suite les colonnes jaunes (résultat, qualification, date, relance).',
    '3. Une ligne n’est jamais « en attente » sans date de relance : les relances échues passent en rouge.',
    '4. « Projets liés » > 1 : le même interlocuteur suit plusieurs projets, groupez-les dans un seul appel.',
    '5. Ouvrez le PDF indiqué en colonne A pour le détail complet de la fiche.',
]
for i, t in enumerate(tips, tips_row + 1):
    wd.cell(i, 2, t).font = font

# ------------------------------------------------------------------ Mode d'emploi
wm = wb.create_sheet("Mode d'emploi", 1)
wm.sheet_view.showGridLines = False
wm.column_dimensions['A'].width = 3
wm.column_dimensions['B'].width = 30
wm.column_dimensions['C'].width = 95
r = 2
wm.cell(r, 2, "Mode d'emploi du fichier de prospection").font = title_font
r += 2
sections = [
    ('Onglets', [
        ('Tableau de bord', 'Indicateurs calculés automatiquement (avancement, joignabilité, conversion, relances échues).'),
        ('Appels', 'La liste d’appels : une ligne par contact et par projet. C’est ici que vous travaillez.'),
        ('Projets', 'Une ligne par projet (= une page de PDF), avec toutes les données extraites de la fiche.'),
        ('Listes', 'Valeurs des listes déroulantes et seuils de priorité (modifiables).'),
    ]),
    ('Code couleur', [
        ('En-tête vert pétrole', 'Données extraites des PDF (ne pas modifier, sauf correction).'),
        ('En-tête orange + cellules jaunes', 'Colonnes à remplir pendant le phoning.'),
        ('En-tête gris', 'Détails complémentaires — colonnes groupées : cliquez sur « − » au-dessus pour les masquer.'),
        ('Ligne verte', 'Prospect intéressé, RDV fixé ou offre demandée.'),
        ('Ligne grisée', 'Dossier clos (pas intéressé, hors cible, déjà équipé, projet terminé).'),
        ('Date de relance en rouge', 'Relance échue sur un dossier encore ouvert : à traiter en priorité.'),
        ('« À rechercher » (orange)', 'Numéro absent de la fiche : chercher sur local.ch / search.ch / site web.'),
    ]),
    ('Colonnes à remplir', [
        ('Résultat dernier appel', 'Ce qui s’est passé : joint (bon interlocuteur ou barrage), messagerie, pas de réponse, occupé, numéro erroné.'),
        ('Qualification', 'Où en est le prospect : à appeler → à rappeler → intéressé → RDV / offre ; ou clos.'),
        ('Nb tentatives', 'Incrémentez à chaque appel (bonne pratique : 5 à 6 tentatives max. avant de clore).'),
        ('Date dernier appel', 'Date du dernier appel (Ctrl + ; insère la date du jour).'),
        ('Prochaine action / Date de relance', 'Toujours une action et une date, sauf dossier clos.'),
        ('Commentaires', 'Nom et fonction de la personne jointe, besoin exprimé, objection, meilleur moment pour rappeler.'),
    ]),
    ('Exemple de ligne remplie', [
        ('Résultat dernier appel', 'Joint – barrage / standard'),
        ('Qualification', 'À rappeler'),
        ('Nb tentatives', '2'),
        ('Date dernier appel', '14/10/2026'),
        ('Prochaine action', 'Rappeler'),
        ('Date de relance', '16/10/2026'),
        ('Commentaires', 'Standard : M. Franco (gérant technique) en visite chantier, rappeler jeudi après 14h. Projet 3e étage confirmé.'),
    ]),
    ('Glossaire (fiches en allemand)', [
        ("Maître d'ouvrage (Bauherr)", 'Le client final / propriétaire ou locataire qui fait réaliser les travaux.'),
        ('Architecte / planificateur', 'Le bureau qui conçoit le projet et prescrit souvent les fournisseurs.'),
        ('Demande de permis déposée', 'Baugesuch eingereicht : projet en phase d’autorisation, travaux à venir (le plus tôt pour se positionner).'),
        ('Permis délivré', 'Baubewilligung erteilt : autorisation obtenue, démarrage des travaux proche (urgent).'),
        ('Réf. projet Infopro', 'Objektnummer : identifiant unique du projet chez Infopro Digital (Baublatt).'),
    ]),
    ('Notes sur les données', [
        ('Source', 'Extraction automatique des 91 PDF du dossier « harworth » (fiches Infopro Digital Schweiz).'),
        ('Fichiers multipages', 'Certains PDF contiennent plusieurs projets : la colonne « Page » indique la page du projet.'),
        ('Doublons', '3 projets figurent dans deux fichiers différents (colonne « Doublon ») : ne les appeler qu’une fois.'),
        ('Liste complémentaire', '5 contacts proviennent du tableau résumé en bas de fiche (nom, ville, téléphone uniquement).'),
        ('Montants', 'Montant des travaux tel qu’annoncé par Infopro (estimation, converti depuis « Mio CHF »).'),
    ]),
]
for title, items in sections:
    wm.cell(r, 2, title).font = h2
    r += 1
    for k, v in items:
        a = wm.cell(r, 2, k)
        b = wm.cell(r, 3, v)
        a.font, b.font = bold, font
        a.alignment = b.alignment = Alignment(wrap_text=True, vertical='top')
        if title == 'Exemple de ligne remplie':
            b.fill = FILL_IN
        for c in (a, b):
            c.border = Border(bottom=thin)
        r += 1
    r += 1

# ------------------------------------------------------------------ workbook-wide
for sh in wb.worksheets:
    sh.sheet_properties.tabColor = {'Tableau de bord': '0E3A40', 'Appels': 'B4541A', 'Projets': '1C5F68'}.get(sh.title, '9AA5A6')
ws.sheet_view.zoomScale = 90
wp.sheet_view.zoomScale = 90
wb.active = wb.sheetnames.index('Appels')
for sh in wb.worksheets:
    sh.sheet_view.tabSelected = sh.title == 'Appels'
ws.print_title_rows = '1:1'
for sh in (wd, wm):
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print('rows', LAST - 1, 'projects', PLAST - 1)
